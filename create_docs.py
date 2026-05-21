from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor as PPTXColor
from pptx.enum.text import PP_ALIGN
import os

OUTPUT_DIR = r"C:\practise\python\060526\docs"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ─────────────────────────────────────────────
# WORD DOCUMENT
# ─────────────────────────────────────────────

def add_heading(doc, text, level=1, color=RGBColor(0x1F, 0x49, 0x7D)):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = color
    return p

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light List Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.font.bold = True
    for r, row in enumerate(rows):
        cells = table.rows[r + 1].cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph()

doc = Document()

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("ML Supervised Learning — Customer Churn Prediction")
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("High-Level Project Document  |  May 2026").font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.add_paragraph()

# 1. Overview
add_heading(doc, "1. Project Overview")
doc.add_paragraph(
    "This project implements an end-to-end supervised machine learning pipeline to predict "
    "customer churn. It covers data generation, model training, REST API deployment via FastAPI, "
    "and automated testing. The solution is production-ready with a clean separation of concerns "
    "across data, model, API, and test layers."
)

# 2. Objectives
add_heading(doc, "2. Objectives")
for obj in [
    "Train a classification model to predict whether a customer will churn (1) or stay (0).",
    "Expose the model as a REST API so any application can consume predictions.",
    "Ensure correctness through an automated test suite.",
    "Structure the codebase for maintainability and easy retraining.",
]:
    doc.add_paragraph(obj, style="List Bullet")

# 3. Architecture
add_heading(doc, "3. System Architecture")
doc.add_paragraph(
    "The system follows a three-layer architecture:"
)
for layer, desc in [
    ("Data Layer", "Generates and stores labeled customer data (customer_churn.csv)."),
    ("Model Layer", "Trains a RandomForestClassifier and persists it as model.pkl."),
    ("API Layer", "FastAPI app loads the model and serves predictions over HTTP."),
]:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(f"{layer}: ")
    run.bold = True
    p.add_run(desc)

doc.add_paragraph()
doc.add_paragraph("Request Flow:", style="Intense Quote")
doc.add_paragraph("Client  →  POST /predict  →  FastAPI  →  model_loader  →  RandomForest  →  JSON response")

# 4. Project Structure
add_heading(doc, "4. Project Structure")
add_table(doc,
    ["Path", "Purpose"],
    [
        ("app/main.py",              "FastAPI app — endpoints and request handling"),
        ("app/schemas.py",           "Pydantic models for input validation and response shape"),
        ("app/model_loader.py",      "Lazy singleton loader for model.pkl"),
        ("day1_ml/project_1/train.py", "Model training script"),
        ("day1_ml/data/data.py",     "Sample data generator (1 000 rows)"),
        ("test/test_api.py",         "Automated API tests using FastAPI TestClient"),
        ("ml/linear_reg/linear_reg.py", "Linear regression reference implementation"),
        (".gitignore",               "Excludes .venv, *.pkl, *.csv, __pycache__"),
    ]
)

# 5. ML Model
add_heading(doc, "5. Machine Learning Model")
doc.add_paragraph(
    "Algorithm: Random Forest Classifier (scikit-learn)\n"
    "Training split: 80% train / 20% test\n"
    "Random state: 42 (reproducible)\n"
    "Estimators: 100 trees"
)
add_heading(doc, "5.1 Features", level=2)
add_table(doc,
    ["Feature", "Type", "Description"],
    [
        ("age",             "Integer",  "Customer age (18–70)"),
        ("contract_type",   "Integer",  "0 = monthly, 1 = yearly"),
        ("monthly_charges", "Float",    "Monthly bill amount (£50–£500)"),
        ("support_calls",   "Integer",  "Number of support calls made (0–10)"),
    ]
)
add_heading(doc, "5.2 Target", level=2)
doc.add_paragraph("churned — Binary label: 0 = customer stayed, 1 = customer churned.")

# 6. API Endpoints
add_heading(doc, "6. API Endpoints")
add_table(doc,
    ["Method", "Endpoint", "Description", "Response"],
    [
        ("GET",  "/health",  "Liveness check — confirms app and model are ready", '{"status": "healthy", "model_loaded": true}'),
        ("POST", "/predict", "Accepts customer features, returns churn prediction", '{"prediction": 0|1, "probability": float, "model_version": "1.0.0"}'),
        ("GET",  "/docs",    "Interactive Swagger UI (auto-generated by FastAPI)", "Browser UI"),
    ]
)

# 7. Testing
add_heading(doc, "7. Testing")
doc.add_paragraph(
    "Tests are located in test/test_api.py and use FastAPI's built-in TestClient "
    "so no running server is needed. Two test cases are included:"
)
add_table(doc,
    ["Test", "What it verifies"],
    [
        ("test_health_endpoint", "GET /health returns HTTP 200"),
        ("test_predict_endpoint", "POST /predict returns HTTP 200 and prediction is 0 or 1"),
    ]
)
doc.add_paragraph("Run tests with:  python -m pytest test/test_api.py -v", style="Intense Quote")

# 8. How to Run
add_heading(doc, "8. How to Run")
steps = [
    ("Generate data",   "python day1_ml/data/data.py"),
    ("Train model",     "python day1_ml/project_1/train.py"),
    ("Start API server","uvicorn app.main:app --host 0.0.0.0 --port 8000"),
    ("Run tests",       "python -m pytest test/test_api.py -v"),
    ("Open API docs",   "http://localhost:8000/docs"),
]
for step, cmd in steps:
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(f"{step}: ")
    run.bold = True
    p.add_run(cmd)

# 9. Tech Stack
add_heading(doc, "9. Technology Stack")
add_table(doc,
    ["Component", "Technology", "Version"],
    [
        ("Language",      "Python",           "3.11"),
        ("ML Framework",  "scikit-learn",     "Latest"),
        ("API Framework", "FastAPI",           "Latest"),
        ("Server",        "Uvicorn",           "Latest"),
        ("Validation",    "Pydantic v2",       "Latest"),
        ("Testing",       "pytest + TestClient","Latest"),
        ("Version Control","Git / GitHub",     "—"),
    ]
)

# 10. Future Improvements
add_heading(doc, "10. Future Improvements")
for item in [
    "Add a retraining endpoint (POST /retrain) that accepts new CSV data.",
    "Integrate MLflow for experiment tracking and model versioning.",
    "Add CI/CD pipeline (GitHub Actions) to auto-test on every push.",
    "Deploy to cloud (AWS / GCP / Azure) behind a load balancer.",
    "Add authentication (API keys or JWT) to the prediction endpoint.",
]:
    doc.add_paragraph(item, style="List Bullet")

word_path = os.path.join(OUTPUT_DIR, "Project_Document.docx")
doc.save(word_path)
print(f"Word document saved: {word_path}")


# ─────────────────────────────────────────────
# POWERPOINT PRESENTATION
# ─────────────────────────────────────────────

DARK_BLUE  = PPTXColor(0x1F, 0x49, 0x7D)
MID_BLUE   = PPTXColor(0x27, 0x80, 0xC2)
WHITE      = PPTXColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = PPTXColor(0xF2, 0xF2, 0xF2)
ACCENT     = PPTXColor(0x00, 0xB0, 0xF0)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

BLANK = 6  # blank layout index

def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[BLANK])

def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def box(slide, left, top, width, height, text, font_size=20,
        bold=False, color=WHITE, align=PP_ALIGN.LEFT,
        bg_color=None, italic=False):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if bg_color:
        txBox.fill.solid()
        txBox.fill.fore_color.rgb = bg_color
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox

def bullet_box(slide, left, top, width, height, items, font_size=16,
               color=WHITE, bullet="•  "):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = bullet + item
        run.font.size = Pt(font_size)
        run.font.color.rgb = color

def accent_bar(slide, top=1.15):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(top), Inches(13.33), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

# ── Slide 1: Title ──────────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, DARK_BLUE)
bar = sl.shapes.add_shape(1, Inches(0), Inches(3.0), Inches(13.33), Inches(2.1))
bar.fill.solid(); bar.fill.fore_color.rgb = MID_BLUE; bar.line.fill.background()
box(sl, 0.5, 3.1, 12.3, 1.0,
    "ML Supervised Learning", 40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
box(sl, 0.5, 4.1, 12.3, 0.6,
    "Customer Churn Prediction  |  FastAPI  |  RandomForest", 20,
    color=ACCENT, align=PP_ALIGN.CENTER)
box(sl, 0.5, 6.5, 12.3, 0.5,
    "May 2026", 14, color=PPTXColor(0xCC, 0xCC, 0xCC), align=PP_ALIGN.CENTER)

# ── Slide 2: Agenda ──────────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, LIGHT_GRAY)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "Agenda", 28, bold=True, color=DARK_BLUE)
items = [
    "Project Overview & Objectives",
    "System Architecture",
    "Machine Learning Model",
    "API Endpoints",
    "Testing Strategy",
    "Technology Stack",
    "How to Run",
    "Future Improvements",
]
bullet_box(sl, 1.0, 1.3, 11, 5.8, items, font_size=18, color=DARK_BLUE, bullet="➤  ")

# ── Slide 3: Overview ────────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, DARK_BLUE)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "Project Overview", 28, bold=True, color=WHITE)
box(sl, 0.5, 1.3, 12, 1.4,
    "An end-to-end supervised ML pipeline that predicts whether a telecom customer "
    "will churn, deployed as a production-ready REST API.",
    18, color=ACCENT)
bullet_box(sl, 0.8, 2.9, 11.5, 4.0, [
    "Generate labeled training data (1,000 customer records)",
    "Train a RandomForestClassifier with scikit-learn",
    "Serve predictions via FastAPI (JSON REST API)",
    "Validate correctness with an automated pytest suite",
    "Version-controlled on GitHub with clean .gitignore",
], font_size=17, color=WHITE)

# ── Slide 4: Architecture ────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, LIGHT_GRAY)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "System Architecture", 28, bold=True, color=DARK_BLUE)

layers = [
    ("Data Layer",  "data.py generates customer_churn.csv\n1,000 rows  |  4 features  |  1 binary label"),
    ("Model Layer", "train.py trains RandomForestClassifier\nSaves serialised model as model.pkl"),
    ("API Layer",   "FastAPI loads model via model_loader.py\nExposes /health and /predict endpoints"),
    ("Test Layer",  "pytest + FastAPI TestClient\nRuns in-process — no live server needed"),
]
colors = [MID_BLUE, DARK_BLUE, MID_BLUE, DARK_BLUE]
for i, (title, desc) in enumerate(layers):
    left = 0.3 + i * 3.25
    card = sl.shapes.add_shape(1, Inches(left), Inches(1.5), Inches(3.0), Inches(2.4))
    card.fill.solid(); card.fill.fore_color.rgb = colors[i]; card.line.fill.background()
    box(sl, left + 0.1, 1.55, 2.8, 0.5, title, 14, bold=True, color=WHITE)
    box(sl, left + 0.1, 2.1,  2.8, 1.6, desc,  12, color=LIGHT_GRAY)

box(sl, 0.5, 4.1, 12, 0.4,
    "Client  →  POST /predict  →  FastAPI  →  model_loader  →  RandomForest  →  JSON",
    14, color=DARK_BLUE, align=PP_ALIGN.CENTER, italic=True)

# ── Slide 5: ML Model ────────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, DARK_BLUE)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "Machine Learning Model", 28, bold=True, color=WHITE)

# Features table (manual)
headers = ["Feature", "Type", "Range"]
rows = [
    ("age",             "int",   "18 – 70"),
    ("contract_type",   "int",   "0 = monthly, 1 = yearly"),
    ("monthly_charges", "float", "£50 – £500"),
    ("support_calls",   "int",   "0 – 10"),
]
col_widths = [2.8, 2.0, 3.5]
starts = [0.5, 3.4, 5.5]
row_h = 0.42
header_top = 1.3
for ci, (hdr, w) in enumerate(zip(headers, col_widths)):
    cell = sl.shapes.add_shape(1, Inches(starts[ci]), Inches(header_top),
                                Inches(w), Inches(row_h))
    cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT; cell.line.fill.background()
    box(sl, starts[ci]+0.05, header_top+0.04, w-0.1, row_h-0.08, hdr, 13, bold=True, color=DARK_BLUE)
for ri, row in enumerate(rows):
    top = header_top + row_h + ri * row_h
    rc = MID_BLUE if ri % 2 == 0 else DARK_BLUE
    for ci, (val, w) in enumerate(zip(row, col_widths)):
        cell = sl.shapes.add_shape(1, Inches(starts[ci]), Inches(top),
                                    Inches(w), Inches(row_h))
        cell.fill.solid(); cell.fill.fore_color.rgb = rc; cell.line.fill.background()
        box(sl, starts[ci]+0.05, top+0.04, w-0.1, row_h-0.08, val, 12, color=WHITE)

box(sl, 9.2, 1.3, 3.8, 0.4, "Algorithm", 13, bold=True, color=ACCENT)
bullet_box(sl, 9.2, 1.7, 3.8, 3.5, [
    "RandomForestClassifier",
    "100 estimators",
    "80 / 20 train-test split",
    "random_state = 42",
    "Output: 0 (stay) / 1 (churn)",
    "Probability score included",
], font_size=13, color=WHITE)

# ── Slide 6: API Endpoints ───────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, LIGHT_GRAY)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "API Endpoints", 28, bold=True, color=DARK_BLUE)

endpoints = [
    ("GET",  "/health",  "Liveness check",       "200 OK  |  {status, model_loaded}"),
    ("POST", "/predict", "Churn prediction",      "200 OK  |  {prediction, probability, model_version}"),
    ("GET",  "/docs",    "Swagger UI (auto-gen)", "Interactive browser docs"),
]
method_colors = {
    "GET":  PPTXColor(0x27, 0xAE, 0x60),
    "POST": PPTXColor(0x27, 0x80, 0xC2),
}
for i, (method, path, desc, resp) in enumerate(endpoints):
    top = 1.5 + i * 1.6
    badge = sl.shapes.add_shape(1, Inches(0.4), Inches(top), Inches(1.1), Inches(0.4))
    badge.fill.solid(); badge.fill.fore_color.rgb = method_colors.get(method, MID_BLUE)
    badge.line.fill.background()
    box(sl, 0.42, top+0.02, 1.06, 0.36, method, 13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    box(sl, 1.6,  top,       3.5,  0.4,  path,   16, bold=True, color=DARK_BLUE)
    box(sl, 5.2,  top,       4.0,  0.4,  desc,   14, color=PPTXColor(0x44,0x44,0x44))
    box(sl, 0.6,  top+0.45,  12.3, 0.7,  resp,   12, color=PPTXColor(0x55,0x55,0x55), italic=True)

# ── Slide 7: Testing ─────────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, DARK_BLUE)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "Testing Strategy", 28, bold=True, color=WHITE)
bullet_box(sl, 0.8, 1.3, 11.5, 2.5, [
    "Framework: pytest + FastAPI TestClient (no live server needed)",
    "test_health_endpoint — asserts GET /health returns HTTP 200",
    "test_predict_endpoint — asserts POST /predict returns 200 and prediction in {0, 1}",
], font_size=17, color=WHITE)
box(sl, 0.5, 3.6, 12, 0.4, "Run command:", 14, bold=True, color=ACCENT)
code_box = sl.shapes.add_textbox(Inches(0.5), Inches(4.1), Inches(12), Inches(0.7))
code_box.fill.solid(); code_box.fill.fore_color.rgb = PPTXColor(0x0D, 0x0D, 0x0D)
tf = code_box.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "  python -m pytest test/test_api.py -v"
run.font.size = Pt(16)
run.font.color.rgb = PPTXColor(0x00, 0xFF, 0x80)
run.font.bold = True

# ── Slide 8: Tech Stack ──────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, LIGHT_GRAY)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "Technology Stack", 28, bold=True, color=DARK_BLUE)
stack = [
    ("Python 3.11",      "Core language"),
    ("scikit-learn",     "ML model training"),
    ("FastAPI",          "REST API framework"),
    ("Uvicorn",          "ASGI server"),
    ("Pydantic v2",      "Request / response validation"),
    ("pytest",           "Automated testing"),
    ("Git / GitHub",     "Version control & hosting"),
]
col = 0
for i, (tech, role) in enumerate(stack):
    row = i % 4
    if i == 4: col = 1
    left = 0.5 + col * 6.4
    top  = 1.4 + row * 1.35
    card = sl.shapes.add_shape(1, Inches(left), Inches(top), Inches(5.9), Inches(1.1))
    card.fill.solid()
    card.fill.fore_color.rgb = DARK_BLUE if col == 0 else MID_BLUE
    card.line.fill.background()
    box(sl, left+0.15, top+0.08, 3.5, 0.45, tech, 15, bold=True, color=WHITE)
    box(sl, left+0.15, top+0.55, 5.5, 0.45, role, 13, color=LIGHT_GRAY)

# ── Slide 9: How to Run ──────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, DARK_BLUE)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "How to Run", 28, bold=True, color=WHITE)
steps = [
    ("1", "Generate Data",    "python day1_ml/data/data.py"),
    ("2", "Train Model",      "python day1_ml/project_1/train.py"),
    ("3", "Start API Server", "uvicorn app.main:app --host 0.0.0.0 --port 8000"),
    ("4", "Run Tests",        "python -m pytest test/test_api.py -v"),
    ("5", "Open API Docs",    "http://localhost:8000/docs"),
]
for i, (num, label, cmd) in enumerate(steps):
    top = 1.3 + i * 1.1
    circle = sl.shapes.add_shape(9, Inches(0.3), Inches(top), Inches(0.55), Inches(0.55))
    circle.fill.solid(); circle.fill.fore_color.rgb = ACCENT; circle.line.fill.background()
    box(sl, 0.3, top+0.03, 0.55, 0.49, num, 14, bold=True, color=DARK_BLUE, align=PP_ALIGN.CENTER)
    box(sl, 1.0, top,      3.5, 0.45, label, 14, bold=True, color=WHITE)
    code = sl.shapes.add_textbox(Inches(4.8), Inches(top), Inches(8.2), Inches(0.5))
    code.fill.solid(); code.fill.fore_color.rgb = PPTXColor(0x0D, 0x0D, 0x0D)
    tf = code.text_frame
    p = tf.paragraphs[0]
    run = p.add_run(); run.text = "  " + cmd
    run.font.size = Pt(13)
    run.font.color.rgb = PPTXColor(0x00, 0xFF, 0x80)

# ── Slide 10: Future Improvements ───────────────────────────────────────────
sl = new_slide(prs)
bg(sl, LIGHT_GRAY)
accent_bar(sl)
box(sl, 0.5, 0.2, 12, 0.9, "Future Improvements", 28, bold=True, color=DARK_BLUE)
bullet_box(sl, 0.8, 1.3, 11.5, 5.5, [
    "Add POST /retrain endpoint to accept new CSV data and retrain on-demand",
    "Integrate MLflow for experiment tracking, metrics, and model versioning",
    "Build a GitHub Actions CI/CD pipeline — auto-test on every push",
    "Deploy to cloud (AWS / GCP / Azure) behind a load balancer",
    "Add JWT / API-key authentication to the /predict endpoint",
    "Build a Streamlit dashboard for non-technical stakeholders",
], font_size=17, color=DARK_BLUE, bullet="➤  ")

# ── Slide 11: Thank You ──────────────────────────────────────────────────────
sl = new_slide(prs)
bg(sl, DARK_BLUE)
bar2 = sl.shapes.add_shape(1, Inches(0), Inches(3.2), Inches(13.33), Inches(1.6))
bar2.fill.solid(); bar2.fill.fore_color.rgb = MID_BLUE; bar2.line.fill.background()
box(sl, 0, 3.35, 13.33, 1.0, "Thank You", 44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
box(sl, 0, 5.2,  13.33, 0.6,
    "github.com/nikhil2sharma/ML-Supervised_Learning-Test_Linear_Reg",
    16, color=ACCENT, align=PP_ALIGN.CENTER)

ppt_path = os.path.join(OUTPUT_DIR, "Project_Summary.pptx")
prs.save(ppt_path)
print(f"PowerPoint saved: {ppt_path}")
